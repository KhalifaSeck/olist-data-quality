"""
Alertes email via smtplib avec rapport Excel en pièce jointe.
DRY_RUN=false → envoi réel avec pièce jointe Excel.
"""
import os
import smtplib
import logging
from datetime import datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
import io

import warnings
warnings.filterwarnings("ignore", category=UserWarning)

import pandas as pd
from dotenv import load_dotenv
load_dotenv()

log = logging.getLogger(__name__)
DRY_RUN = os.getenv("DRY_RUN", "true").lower() == "true"


def build_excel_attachment(anomalies: list) -> bytes:
    """Génère un fichier Excel avec les anomalies détectées."""
    data = [{
        "Date":        a.date,
        "Métrique":    a.metric,
        "Sévérité":    a.severity.upper(),
        "Valeur":      a.value,
        "Attendu":     a.expected,
        "Z-Score":     a.z_score,
        "Description": a.description
    } for a in anomalies]

    df = pd.DataFrame(data)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Anomalies")

        # Mise en forme
        workbook  = writer.book
        worksheet = writer.sheets["Anomalies"]

        # Largeur des colonnes
        col_widths = {
            "A": 12,  # Date
            "B": 20,  # Métrique
            "C": 12,  # Sévérité
            "D": 12,  # Valeur
            "E": 12,  # Attendu
            "F": 10,  # Z-Score
            "G": 60,  # Description
        }
        for col, width in col_widths.items():
            worksheet.column_dimensions[col].width = width

        # Couleur des lignes selon sévérité
        from openpyxl.styles import PatternFill, Font
        red_fill    = PatternFill("solid", fgColor="FFE5E5")
        yellow_fill = PatternFill("solid", fgColor="FFF8E1")
        bold_font   = Font(bold=True)

        for row in worksheet.iter_rows(min_row=2, max_row=len(df) + 1):
            severity = row[2].value
            fill = red_fill if severity == "CRITICAL" else yellow_fill
            for cell in row:
                cell.fill = fill

        # Header en gras
        for cell in worksheet[1]:
            cell.font = bold_font

    buffer.seek(0)
    return buffer.read()


def build_html_body(anomalies: list) -> str:
    critical = [a for a in anomalies if a.severity == "critical"]
    warnings_ = [a for a in anomalies if a.severity == "warning"]

    rows = ""
    for a in anomalies:
        color = "#e53e3e" if a.severity == "critical" else "#dd6b20"
        rows += f"""
        <tr>
            <td style='padding:8px 12px;border-bottom:1px solid #eee'>{a.date}</td>
            <td style='padding:8px 12px;border-bottom:1px solid #eee'>{a.metric}</td>
            <td style='padding:8px 12px;border-bottom:1px solid #eee;
                color:{color};font-weight:bold'>{a.severity.upper()}</td>
            <td style='padding:8px 12px;border-bottom:1px solid #eee'>{a.description}</td>
        </tr>
        """
    return f"""
    <html>
    <body style='font-family:Arial,sans-serif;color:#333'>
        <h2 style='color:#2d3748'>Olist Data Quality — Rapport d'anomalies</h2>
        <p style='color:#718096'>
            Généré le {datetime.now().strftime('%Y-%m-%d %H:%M')}
        </p>
        <p>
            Total : <b>{len(anomalies)}</b> anomalie(s) |
            🔴 Critical : <b>{len(critical)}</b> |
            🟡 Warning : <b>{len(warnings_)}</b>
        </p>
        <p style='color:#718096;font-size:12px'>
            Le rapport Excel complet est en pièce jointe.
        </p>
        <table style='border-collapse:collapse;width:100%;font-size:13px'>
            <thead>
                <tr style='background:#f7fafc'>
                    <th style='padding:8px 12px;text-align:left'>Date</th>
                    <th style='padding:8px 12px;text-align:left'>Métrique</th>
                    <th style='padding:8px 12px;text-align:left'>Sévérité</th>
                    <th style='padding:8px 12px;text-align:left'>Détail</th>
                </tr>
            </thead>
            <tbody>{rows}</tbody>
        </table>
        <p style='color:#718096;font-size:12px;margin-top:20px'>
            Pipeline Olist Data Quality
        </p>
    </body>
    </html>
    """


def send_alert(anomalies: list,
               subject: str = "Olist DQ — Anomalies détectées"):
    if not anomalies:
        log.info("No anomalies — no email sent.")
        return

    sender   = os.getenv("SMTP_SENDER")
    receiver = os.getenv("SMTP_RECEIVER")

    if not sender or not receiver:
        log.warning("SMTP_SENDER ou SMTP_RECEIVER non défini — skipped.")
        return

    msg = MIMEMultipart("mixed")
    msg["Subject"] = subject
    msg["From"]    = sender
    msg["To"]      = receiver

    # Corps HTML
    msg.attach(MIMEText(build_html_body(anomalies), "html"))

    # Pièce jointe Excel
    excel_data = build_excel_attachment(anomalies)
    filename   = f"olist_anomalies_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    part = MIMEBase("application",
                    "vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    part.set_payload(excel_data)
    encoders.encode_base64(part)
    part.add_header("Content-Disposition",
                    f"attachment; filename={filename}")
    msg.attach(part)

    if DRY_RUN:
        log.info(f"[DRY RUN] Email simulé → {receiver}")
        log.info(f"[DRY RUN] Pièce jointe : {filename}")
        _print_summary(anomalies)
        return

    try:
        with smtplib.SMTP(
            os.getenv("SMTP_HOST", "smtp.gmail.com"),
            int(os.getenv("SMTP_PORT", 587))
        ) as server:
            server.starttls()
            server.login(
                os.getenv("SMTP_USER"),
                os.getenv("SMTP_PASSWORD")
            )
            server.sendmail(sender, receiver, msg.as_string())
        log.info(f"Email + Excel envoyé à {receiver} ({len(anomalies)} anomalie(s))")
    except Exception as e:
        log.error(f"Erreur envoi email : {e}")


def _print_summary(anomalies: list):
    critical = [a for a in anomalies if a.severity == "critical"]
    warnings_ = [a for a in anomalies if a.severity == "warning"]
    print("\n" + "="*60)
    print(f"  RAPPORT — {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    print("="*60)
    print(f"  Total    : {len(anomalies)} anomalie(s)")
    print(f"  Critical : {len(critical)}")
    print(f"  Warning  : {len(warnings_)}")
    print("-"*60)
    if critical:
        print("  CRITIQUES :")
        for a in critical:
            print(f"    [{a.date}] {a.description}")
    if warnings_:
        print("  WARNINGS (5 premiers) :")
        for a in warnings_[:5]:
            print(f"    [{a.date}] {a.description}")
    print("="*60 + "\n")


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    import sys
    sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'monitoring'))
    from anomaly_detection import run_all_checks
    anomalies = run_all_checks()
    send_alert(anomalies)